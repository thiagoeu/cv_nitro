import io
import os
from typing import Union, BinaryIO
from pypdf import PdfReader
import docx

def extract_text_from_file(file_input: Union[str, BinaryIO, any], filename: str = None) -> str:
    """
    Extrai texto a partir de um arquivo nos formatos PDF, DOCX, MD ou TXT.
    
    file_input pode ser um caminho no sistema de arquivos, um objeto BytesIO 
    ou um UploadedFile do Streamlit.
    """
    if filename is None:
        if isinstance(file_input, str):
            filename = os.path.basename(file_input)
        elif hasattr(file_input, "name"):
            filename = file_input.name
        else:
            filename = ""

    ext = os.path.splitext(filename)[1].lower()

    # Prepara o stream em memória se for caminho de arquivo
    if isinstance(file_input, str):
        with open(file_input, "rb") as f:
            stream = io.BytesIO(f.read())
    elif hasattr(file_input, "read"):
        content = file_input.read()
        # Se for UploadedFile do Streamlit, redefinir ponteiro se possível
        if hasattr(file_input, "seek"):
            file_input.seek(0)
        stream = io.BytesIO(content)
    else:
        raise ValueError("Entrada de arquivo inválida.")

    extracted_text = ""

    if ext == ".pdf":
        try:
            reader = PdfReader(stream)
            pages_text = []
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    pages_text.append(text)
            extracted_text = "\n\n".join(pages_text)
        except Exception as e:
            raise RuntimeError(f"Erro ao extrair texto do PDF ({filename}): {e}")

    elif ext == ".docx":
        try:
            doc = docx.Document(stream)
            content_parts = []
            
            # Extrai parágrafos normais
            for p in doc.paragraphs:
                p_text = p.text.strip()
                if p_text:
                    content_parts.append(p_text)
                    
            # Extrai texto de tabelas (frequentes em layouts de currículo)
            for table in doc.tables:
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    unique_cells = []
                    for c in row_cells:
                        if not unique_cells or c != unique_cells[-1]:
                            unique_cells.append(c)
                    if unique_cells:
                        content_parts.append(" | ".join(unique_cells))
                        
            extracted_text = "\n\n".join(content_parts)
        except Exception as e:
            raise RuntimeError(f"Erro ao extrair texto do DOCX ({filename}): {e}")

    elif ext in [".md", ".markdown", ".txt", ""]:
        try:
            raw = stream.getvalue()
            try:
                extracted_text = raw.decode("utf-8")
            except UnicodeDecodeError:
                extracted_text = raw.decode("latin-1")
        except Exception as e:
            raise RuntimeError(f"Erro ao extrair texto do arquivo ({filename}): {e}")

    elif ext == ".doc":
        raise ValueError("O formato legado .doc (Microsoft Word 97-2003) não é suportado. Por favor, salve ou converta o arquivo para .docx ou .pdf.")

    else:
        raise ValueError(f"Formato de arquivo não suportado: {ext}. Formatos suportados: .pdf, .docx, .md, .txt")

    return extracted_text.strip()
