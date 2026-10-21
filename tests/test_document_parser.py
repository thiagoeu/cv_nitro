import os
import io
import pytest
from src.core.document_parser import extract_text_from_file

def test_extract_text_markdown_txt():
    content = "# Curriculo\nDesenvolvedor Python com 5 anos de experiência."
    stream = io.BytesIO(content.encode("utf-8"))
    stream.name = "curriculo.md"
    extracted = extract_text_from_file(stream)
    assert "Desenvolvedor Python" in extracted
    assert "# Curriculo" in extracted

def test_extract_text_docx():
    import docx
    doc = docx.Document()
    doc.add_heading("Currículo João Silva", 0)
    doc.add_paragraph("Engenheiro de Software com foco em Python e IA.")
    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)
    stream.name = "curriculo.docx"
    
    extracted = extract_text_from_file(stream)
    assert "Currículo João Silva" in extracted
    assert "Engenheiro de Software" in extracted

def test_extract_text_pdf():
    from pypdf import PdfWriter
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    stream = io.BytesIO()
    writer.write(stream)
    stream.seek(0)
    stream.name = "curriculo.pdf"
    
    # Testa que não quebra ao ler pdf (mesmo que página em branco retorne vazio)
    extracted = extract_text_from_file(stream)
    assert isinstance(extracted, str)

def test_extract_text_docx_with_tables():
    import docx
    doc = docx.Document()
    doc.add_heading("Currículo Com Tabelas", 0)
    doc.add_paragraph("Resumo do candidato.")
    
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Habilidades Backend"
    table.cell(0, 1).text = "Python, FastAPI, Docker"
    table.cell(1, 0).text = "Banco de Dados"
    table.cell(1, 1).text = "PostgreSQL, MongoDB"
    
    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)
    stream.name = "curriculo_tabela.docx"
    
    extracted = extract_text_from_file(stream)
    assert "Habilidades Backend" in extracted
    assert "Python, FastAPI, Docker" in extracted
    assert "PostgreSQL, MongoDB" in extracted

def test_extract_text_unsupported_formats():
    stream_doc = io.BytesIO(b"fake doc content")
    stream_doc.name = "antigo.doc"
    with pytest.raises(ValueError, match="legado .doc"):
        extract_text_from_file(stream_doc)

    stream_unknown = io.BytesIO(b"binary")
    stream_unknown.name = "arquivo.exe"
    with pytest.raises(ValueError, match="Formato de arquivo não suportado"):
        extract_text_from_file(stream_unknown)


def test_extract_text_docx_keeps_table_in_document_order():
    import docx
    doc = docx.Document()
    doc.add_paragraph("HABILIDADES TÉCNICAS")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Back-end:"
    table.cell(0, 1).text = "Node.js, NestJS"
    doc.add_paragraph("EXPERIÊNCIA PROFISSIONAL")

    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)
    stream.name = "ordem.docx"

    extracted = extract_text_from_file(stream)
    assert extracted == "HABILIDADES TÉCNICAS\n\n| Back-end: | Node.js, NestJS |\n\nEXPERIÊNCIA PROFISSIONAL"
